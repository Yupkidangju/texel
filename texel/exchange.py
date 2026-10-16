from __future__ import annotations

import csv
import io
import os
import sqlite3
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError

from .addresses import CellRange, column_name, parse_address
from .formula import Evaluator
from .model import MAX_CELLS, BorderEdge, CellStyle, Validation, Workbook
from .values import exact_text

XLSX_NOTICE = "XLSX: cells/formulas/basic styles/notes/sheets supported; VBA, external links and Excel-only objects are not preserved."


def safe_path(root: Path | str, path: Path | str) -> Path:
    root = Path(root).resolve()
    path = Path(path)
    result = (root / path if not path.is_absolute() else path).resolve()
    if not result.is_relative_to(root):
        raise ValueError("Path is outside the configured MCP root")
    return result


def load_file(path: Path | str, encoding="auto") -> Workbook:
    try:
        return _load_file(path, encoding)
    except (sqlite3.DatabaseError, zipfile.BadZipFile, csv.Error, ParseError, KeyError) as exc:
        raise ValueError(f"Invalid spreadsheet file: {exc}") from exc


def _load_file(path: Path | str, encoding="auto") -> Workbook:
    path = Path(path)
    if not path.is_file() or path.stat().st_size > 64 * 1024 * 1024:
        raise ValueError("Input must be an existing file of at most 64 MiB")
    suffix = path.suffix.lower()
    if suffix == ".xls":
        from .legacy_excel import load_xls

        return load_xls(path)
    if suffix == ".gsheet":
        from .google_sheets import load_gsheet

        return load_gsheet(path)
    if suffix == ".texel":
        from .store import Store

        return Store(path, read_only=True).load()[0]
    if suffix in {".csv", ".tsv"}:
        book, count = Workbook(), 0
        csv.field_size_limit(32767)
        if encoding not in {"auto", "utf-8", "utf-8-sig", "utf-16", "cp949"}:
            raise ValueError("CSV encoding must be auto/utf-8-sig/utf-16/cp949")
        data = path.read_bytes()
        if encoding == "auto":
            if data.startswith((b"\xff\xfe", b"\xfe\xff")):
                decoded = data.decode("utf-16")
            else:
                try:
                    decoded = data.decode("utf-8-sig")
                except UnicodeDecodeError:
                    decoded = data.decode("cp949")
        else:
            decoded = data.decode(encoding)
        with io.StringIO(decoded, newline="") as stream:
            for r, row in enumerate(
                csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ",")
            ):
                for c, value in enumerate(row):
                    if value:
                        count += 1
                        if count > MAX_CELLS:
                            raise ValueError("CSV exceeds stored cell limit")
                        book.sheet().set_raw(r, c, value)
        book.validate()
        return book
    if suffix != ".xlsx":
        raise ValueError("Supported files: .texel, .csv, .tsv, .xlsx, .xls, .gsheet")
    return load_xlsx(path)


def load_xlsx(path):
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 128 * 1024 * 1024:
            raise ValueError("XLSX uncompressed size exceeds 128 MiB")
    source = load_workbook(path, data_only=False, keep_links=False)
    book, count = Workbook(sheets=[]), 0
    try:
        for ws in source.worksheets:
            if ws.max_row * ws.max_column > 1000000:
                raise ValueError("XLSX worksheet rectangle exceeds 1,000,000 cells")
            index = book.add_sheet(ws.title)
            target = book.sheet(index)
            for row in ws.iter_rows():
                for original in row:
                    if (
                        original.value is None
                        and not original.has_style
                        and original.comment is None
                    ):
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("XLSX exceeds stored cell limit")
                    r, c = original.row - 1, original.column - 1
                    value = original.value
                    if isinstance(value, datetime):
                        raw = (
                            value.date().isoformat()
                            if value.time().isoformat() == "00:00:00"
                            else value.isoformat()
                        )
                    elif isinstance(value, date):
                        raw = value.isoformat()
                    elif isinstance(value, bool):
                        raw = "TRUE" if value else "FALSE"
                    elif value is None:
                        raw = ""
                    elif original.data_type == "f":
                        raw = str(value)
                    elif isinstance(value, str):
                        raw = "'" + value
                    else:
                        raw = str(value)
                    target.set_raw(r, c, raw)
                    cell = target.cell(r, c)
                    fmt = original.number_format
                    style = CellStyle(
                        bold=bool(original.font.b),
                        underline=bool(original.font.u),
                        italic=bool(original.font.i),
                        strike=bool(original.font.strike),
                        wrap=bool(original.alignment.wrap_text),
                        vertical_align={"center": "middle", "bottom": "bottom"}.get(
                            original.alignment.vertical, "top"
                        ),
                        align=original.alignment.horizontal
                        if original.alignment.horizontal in {"left", "right", "center"}
                        else "auto",
                    )
                    from openpyxl.styles.colors import COLOR_INDEX

                    for side in ("left", "right", "top", "bottom"):
                        edge = getattr(original.border, side)
                        if edge is not None and edge.style:
                            kind = (
                                edge.style
                                if edge.style
                                in {"thin", "medium", "thick", "double", "dashed", "dotted"}
                                else "dashed"
                                if "dash" in edge.style.lower()
                                else "thin"
                            )
                            color = ""
                            if edge.color is not None:
                                if edge.color.type == "rgb":
                                    color = "#" + edge.color.rgb[-6:]
                                elif edge.color.type == "indexed" and 0 <= edge.color.indexed < len(
                                    COLOR_INDEX
                                ):
                                    color = "#" + COLOR_INDEX[edge.color.indexed][-6:]
                            style.borders[side] = BorderEdge(kind, color)
                    style.number_format = (
                        "percent"
                        if "%" in fmt
                        else "currency"
                        if any(v in fmt for v in ("₩", "$", "€", "£"))
                        else "date"
                        if isinstance(value, (date, datetime))
                        else "number"
                        if fmt != "General" and "0" in fmt
                        else "general"
                    )
                    import re

                    decimals = re.search(r"\.([0#]+)", fmt)
                    style.decimals = min(10, len(decimals[1])) if decimals else 0
                    if original.font.color is not None and original.font.color.type == "rgb":
                        style.foreground = "#" + original.font.color.rgb[-6:]
                    if original.fill.fill_type == "solid" and original.fill.fgColor.type == "rgb":
                        style.background = "#" + original.fill.fgColor.rgb[-6:]
                    cell.style = style
                    cell.note = original.comment.text if original.comment else ""
                    cell.locked = bool(original.protection.locked) if ws.protection.sheet else False
            for key, dimension in ws.column_dimensions.items():
                _, c = parse_address(key + "1")
                if dimension.width:
                    target.column_widths[c] = max(4, min(60, int(dimension.width)))
                if dimension.hidden:
                    target.hidden_columns.add(c)
            target.hidden_rows = {
                int(r) - 1 for r, dimension in ws.row_dimensions.items() if dimension.hidden
            }
            target.row_heights = {
                int(r) - 1: max(1, min(12, round(d.height / 15)))
                for r, d in ws.row_dimensions.items()
                if d.height
            }
            if ws.freeze_panes:
                target.freeze_rows, target.freeze_columns = parse_address(str(ws.freeze_panes))
            target.merges = [str(merge) for merge in ws.merged_cells.ranges]
            target.protected = bool(ws.protection.sheet)
            for rule in ws.data_validations.dataValidation:
                validation = None
                if rule.type == "list" and rule.formula1 and rule.formula1.startswith('"'):
                    validation = Validation(
                        "list", rule.formula1.strip('"').split(","), allow_blank=rule.allowBlank
                    )
                elif rule.type in {"decimal", "whole"} and rule.operator == "between":
                    try:
                        validation = Validation(
                            minimum=float(rule.formula1),
                            maximum=float(rule.formula2),
                            allow_blank=rule.allowBlank,
                        )
                    except ValueError, TypeError:
                        pass
                if validation:
                    for span in str(rule.sqref).split():
                        for r, c in CellRange.parse(span).coordinates(limit=10000):
                            target.cell(r, c).validation = validation
        for name, definition in source.defined_names.items():
            destinations = list(definition.destinations)
            if len(destinations) == 1:
                sheet_name, area = destinations[0]
                try:
                    book.sheet_index(sheet_name)
                    span = str(CellRange.parse(area))
                except ValueError:
                    continue
                book.named_ranges[name] = (sheet_name, span)
    finally:
        source.close()
    book.validate()
    return book


def export_file(book: Workbook, path, sheet=None, formulas=False):
    path = Path(path).resolve()
    suffix = path.suffix.lower()
    if suffix not in {".csv", ".tsv", ".xlsx", ".texel"}:
        raise ValueError("Supported export extensions: .texel, .csv, .tsv, .xlsx")
    if suffix == ".texel" and path.exists():
        raise FileExistsError(
            "Use a new native workbook path; existing live databases cannot be replaced"
        )
    book.validate()
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix=".texel-export-", suffix=suffix, dir=path.parent)
    os.close(handle)
    temporary = Path(temporary)
    try:
        if suffix == ".texel":
            from .session import Session

            temporary.unlink()
            session = Session(temporary)
            session.apply("import", lambda w: replace_workbook(w, book))
            with session.store.connect() as db:
                db.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                db.execute("PRAGMA journal_mode=DELETE")
        elif suffix in {".csv", ".tsv"}:
            active = book.active_sheet if sheet is None else sheet
            target, evaluator = book.sheet(active), Evaluator(book)
            region = target.used_range
            if region.size > 1000000:
                raise ValueError("Export rectangle exceeds 1,000,000 cells")
            with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream, delimiter="\t" if suffix == ".tsv" else ",")
                for r in range(region.bottom + 1):
                    output = []
                    for c in range(region.right + 1):
                        value = evaluator.value(active, r, c)
                        result = target.cell(r, c, False).raw if formulas else exact_text(value)
                        if (
                            not formulas
                            and isinstance(value, str)
                            and result.startswith(("=", "+", "-", "@", "\t", "\r", "\n"))
                        ):
                            result = "'" + result
                        output.append(result)
                    writer.writerow(output)
                stream.flush()
                os.fsync(stream.fileno())
        else:
            save_xlsx(book, temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
        for suffix in ("-wal", "-shm"):
            sidecar = Path(str(temporary) + suffix)
            if sidecar.exists():
                sidecar.unlink()
    return path


def replace_workbook(target, source):
    from copy import deepcopy

    target.sheets = deepcopy(source.sheets)
    target.named_ranges = deepcopy(source.named_ranges)
    target.active_sheet = source.active_sheet


def save_xlsx(book, path):
    from openpyxl import Workbook as ExcelWorkbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    for sheet in book.sheets:
        for original in sheet.cells.values():
            if ILLEGAL_CHARACTERS_RE.search(original.raw) or ILLEGAL_CHARACTERS_RE.search(
                original.note
            ):
                raise ValueError(
                    "XLSX cannot represent these control characters; use native format to preserve the raw data"
                )

    result = ExcelWorkbook()
    result.remove(result.active)
    for sheet in book.sheets:
        ws = result.create_sheet(sheet.name)
        for (r, c), original in sheet.cells.items():
            raw = original.raw
            from .model import literal

            value = raw if raw.startswith("=") else literal(raw)
            cell = ws.cell(r + 1, c + 1, value)
            if raw.startswith("'"):
                cell.data_type = "s"
            style = original.style
            color = style.foreground.lstrip("#") if style.foreground.startswith("#") else None
            cell.font = Font(
                bold=style.bold,
                underline="single" if style.underline else None,
                color=color,
                italic=style.italic,
                strike=style.strike,
            )
            if style.background.startswith("#"):
                cell.fill = PatternFill("solid", fgColor=style.background.lstrip("#"))
            cell.alignment = Alignment(
                horizontal=style.align if style.align != "auto" else None,
                vertical="center" if style.vertical_align == "middle" else style.vertical_align,
                wrap_text=style.wrap,
            )
            cell.border = Border(
                **{
                    side: Side(
                        style=edge.style if edge.style != "none" else None,
                        color=edge.color.lstrip("#") or None,
                    )
                    for side, edge in style.borders.items()
                }
            )
            formats = {
                "general": "General",
                "number": "#,##0" + ("." + "0" * style.decimals if style.decimals else ""),
                "percent": "0" + ("." + "0" * style.decimals if style.decimals else "") + "%",
                "currency": '"₩"#,##0' + ("." + "0" * style.decimals if style.decimals else ""),
                "date": "yyyy-mm-dd",
            }
            cell.number_format = formats[style.number_format]
            if original.note:
                cell.comment = Comment(original.note, "Texel")
            cell.protection = Protection(locked=original.locked)
            if original.validation:
                rule = original.validation
                if rule.kind == "list":
                    validator = DataValidation(
                        type="list",
                        formula1='"' + ",".join(rule.choices).replace('"', '""') + '"',
                        allow_blank=rule.allow_blank,
                    )
                else:
                    low, high = rule.minimum, rule.maximum
                    operator = (
                        "between"
                        if low is not None and high is not None
                        else "greaterThanOrEqual"
                        if low is not None
                        else "lessThanOrEqual"
                    )
                    validator = DataValidation(
                        type="decimal",
                        operator=operator,
                        formula1=str(low if low is not None else high),
                        formula2=str(high) if low is not None and high is not None else None,
                        allow_blank=rule.allow_blank,
                    )
                validator.showErrorMessage = True
                ws.add_data_validation(validator)
                validator.add(cell)
        for c, width in sheet.column_widths.items():
            ws.column_dimensions[column_name(c)].width = width
        for c in sheet.hidden_columns:
            ws.column_dimensions[column_name(c)].hidden = True
        for r in sheet.hidden_rows:
            ws.row_dimensions[r + 1].hidden = True
        for r, height in sheet.row_heights.items():
            ws.row_dimensions[r + 1].height = height * 15
        if sheet.freeze_rows or sheet.freeze_columns:
            ws.freeze_panes = f"{column_name(sheet.freeze_columns)}{sheet.freeze_rows + 1}"
        for merge in sheet.merges:
            ws.merge_cells(merge)
        ws.protection.sheet = sheet.protected
        for rule in sheet.conditional_rules:
            ws.conditional_formatting.add(
                rule["range"],
                CellIsRule(
                    operator=rule.get("operator", "greaterThan"),
                    formula=[str(rule["value"])],
                    fill=PatternFill("solid", fgColor=rule.get("color", "#0f766e").lstrip("#")),
                ),
            )
        for description in sheet.charts:
            span = CellRange.parse(description["range"])
            chart = BarChart() if description.get("kind", "bar") == "bar" else LineChart()
            chart.title = description.get("title", "Texel")
            chart.add_data(
                Reference(
                    ws,
                    min_col=span.left + 2,
                    max_col=span.right + 1,
                    min_row=span.top + 1,
                    max_row=span.bottom + 1,
                ),
                titles_from_data=True,
            )
            chart.set_categories(
                Reference(ws, min_col=span.left + 1, min_row=span.top + 2, max_row=span.bottom + 1)
            )
            ws.add_chart(chart, f"{column_name(min(span.right + 2, 16383))}{span.top + 1}")
    for name, (sheet_name, area) in book.named_ranges.items():
        if "#REF!" not in area:
            span = CellRange.parse(area)
            start = f"${column_name(span.left)}${span.top + 1}"
            end = f"${column_name(span.right)}${span.bottom + 1}"
            result.defined_names.add(
                DefinedName(
                    name, attr_text="'" + sheet_name.replace("'", "''") + "'!" + start + ":" + end
                )
            )
    result.active = book.active_sheet
    result.save(path)
    result.close()
